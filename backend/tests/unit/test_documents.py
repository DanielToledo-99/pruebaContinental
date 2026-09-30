import zipfile
from io import BytesIO

import pytest
from docx import Document as WordDocument
from docx.oxml import parse_xml
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.application.chunking import split_text
from app.domain.exceptions import (
    DocumentTooLarge,
    InvalidDocument,
    UnsupportedDocumentType,
)
from app.infrastructure.document_processing.parsers import (
    DOCX_MIME,
    DocxParser,
    FileDocumentProcessor,
)

NAMESPACES = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
)


def make_pdf(text: str = "El proyecto reduce el consumo de agua.") -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font}),
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode("ascii"))
    page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def make_docx(
    text: str = "El proyecto reduce el consumo de agua.", blocks: tuple[str, ...] = ()
) -> bytes:
    """A Word file with one paragraph plus raw body-level XML blocks."""
    document = WordDocument()
    document.add_paragraph(text)
    body = document.element.body
    for block in blocks:
        body.insert(len(body) - 1, parse_xml(block.replace("NS", NAMESPACES, 1)))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


@pytest.mark.parametrize(
    "name,mime,content,error",
    [
        ("bad.exe", "application/octet-stream", b"data", UnsupportedDocumentType),
        ("doc.pdf", "text/plain", b"%PDF-", UnsupportedDocumentType),
        ("doc.pdf", "application/pdf", b"fake", InvalidDocument),
        ("doc.txt", "text/plain", b"", InvalidDocument),
        ("doc.doc", "application/msword", b"\xd0\xcf\x11\xe0", UnsupportedDocumentType),
        ("doc.docx", "application/pdf", b"PK\x03\x04", UnsupportedDocumentType),
        ("doc.docx", DOCX_MIME, b"not a zip", InvalidDocument),
    ],
)
def test_validate_document(name, mime, content, error):
    with pytest.raises(error):
        FileDocumentProcessor().validate(name, mime, content)


def test_size_and_filename():
    processor = FileDocumentProcessor(max_bytes=4)
    assert processor.validate("../../my\\file.TXT", "text/plain", b"text") == "file.txt"
    with pytest.raises(DocumentTooLarge):
        processor.validate("a.txt", "text/plain", b"12345")


def test_protected_word_files_are_explained():
    with pytest.raises(InvalidDocument, match="protegidos con contraseña"):
        FileDocumentProcessor().validate("doc.docx", DOCX_MIME, b"\xd0\xcf\x11\xe0")


async def test_parse_docx_in_reading_order():
    document = WordDocument()
    document.add_heading("Proyecto Agua Clara")
    document.add_paragraph("El objetivo es ahorrar agua.")
    table = document.add_table(rows=2, cols=3)
    header = ["Concepto", "Detalle", "Monto"]
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = text
    merged = table.cell(1, 0).merge(table.cell(1, 1))
    merged.text = "Presupuesto"
    table.cell(1, 2).text = "15 000 soles"
    body = document.element.body
    for block in (
        # Tracked changes: the insertion counts, the deletion does not.
        "<w:p NS><w:r><w:t xml:space='preserve'>Texto </w:t></w:r>"
        "<w:ins w:id='1' w:author='A'><w:r><w:t>insertado</w:t></w:r></w:ins>"
        "<w:del w:id='2' w:author='A'><w:r><w:delText>borrado</w:delText></w:r>"
        "</w:del><w:r><w:t xml:space='preserve'> y </w:t></w:r>"
        "<w:hyperlink r:id='rId99'><w:r><w:t>enlace</w:t></w:r></w:hyperlink></w:p>",
        "<w:sdt NS><w:sdtContent><w:p><w:r><w:t>Dentro de un control de "
        "contenido.</w:t></w:r></w:p></w:sdtContent></w:sdt>",
        # Text boxes are stored twice; the legacy fallback must not repeat.
        "<w:p NS><w:r><w:t>Ancla</w:t></w:r><w:r><mc:AlternateContent>"
        "<mc:Choice Requires='wps'><w:drawing><w:txbxContent><w:p><w:r>"
        "<w:t>En un cuadro de texto.</w:t></w:r></w:p></w:txbxContent></w:drawing>"
        "</mc:Choice><mc:Fallback><w:pict><w:txbxContent><w:p><w:r>"
        "<w:t>En un cuadro de texto.</w:t></w:r></w:p></w:txbxContent></w:pict>"
        "</mc:Fallback></mc:AlternateContent></w:r></w:p>",
        # Tab stops in paragraph properties are not text.
        "<w:p NS><w:pPr><w:tabs><w:tab w:val='left' w:pos='720'/></w:tabs></w:pPr>"
        "<w:r><w:t>Con</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>tabulador</w:t>"
        "</w:r></w:p>",
    ):
        body.insert(len(body) - 1, parse_xml(block.replace("NS", NAMESPACES, 1)))
    output = BytesIO()
    document.save(output)
    text = await FileDocumentProcessor().extract("doc.docx", output.getvalue())
    assert text.split("\n") == [
        "Proyecto Agua Clara",
        "El objetivo es ahorrar agua.",
        "Concepto | Detalle | Monto",
        "Presupuesto | 15 000 soles",
        "Texto insertado y enlace",
        "Dentro de un control de contenido.",
        "Ancla",
        "En un cuadro de texto.",
        "Con tabulador",
    ]


async def test_docx_limits_and_damaged_files():
    with pytest.raises(DocumentTooLarge):
        DocxParser(max_unzipped=1000).parse(make_docx(), 10_000)
    with pytest.raises(DocumentTooLarge):
        await FileDocumentProcessor(max_chars=5).extract("doc.docx", make_docx())
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w") as other:
        other.writestr("hola.txt", "no es Word")
    with pytest.raises(InvalidDocument, match="Word"):
        await FileDocumentProcessor().extract("doc.docx", archive.getvalue())


async def test_parse_pdf():
    result = await FileDocumentProcessor().extract("doc.pdf", make_pdf())
    assert "reduce el consumo" in result


def encrypt_pdf(content: bytes, user_password: str) -> bytes:
    writer = PdfWriter(clone_from=PdfReader(BytesIO(content)))
    writer.encrypt(
        user_password=user_password, owner_password="owner", algorithm="AES-128"
    )
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


async def test_pdf_with_only_owner_password_is_readable():
    processor = FileDocumentProcessor()
    result = await processor.extract("doc.pdf", encrypt_pdf(make_pdf(), ""))
    assert "reduce el consumo" in result
    with pytest.raises(InvalidDocument, match="contraseña"):
        await processor.extract("doc.pdf", encrypt_pdf(make_pdf(), "secreta"))


async def test_empty_pdf_and_binary_text():
    processor = FileDocumentProcessor()
    with pytest.raises(InvalidDocument):
        await processor.extract("doc.pdf", make_pdf(""))
    with pytest.raises(InvalidDocument):
        await processor.extract("doc.txt", b"abc\x00def")


async def test_extracted_text_limit():
    with pytest.raises(DocumentTooLarge):
        await FileDocumentProcessor(max_chars=5).extract("doc.pdf", make_pdf())


def test_chunks_preserve_text_and_have_overlap():
    text = "".join(str(i % 10) for i in range(2500))
    chunks = split_text(text, 800, 100)
    assert all(0 < len(chunk) <= 800 for chunk in chunks)
    assert chunks[0][-100:] == chunks[1][:100]
    reconstructed = chunks[0] + "".join(chunk[100:] for chunk in chunks[1:])
    assert reconstructed == text


async def test_extraction_keeps_line_and_paragraph_breaks():
    content = "Título\r\n\r\n\r\nPrimera  línea\t \ncontinúa.".encode()
    text = await FileDocumentProcessor().extract("doc.txt", content)
    assert text == "Título\n\nPrimera línea\ncontinúa."


def test_chunks_prefer_paragraph_and_sentence_boundaries():
    paragraph = ("Primera oración del párrafo. " * 20).strip()
    assert split_text(f"{paragraph}\n\n{paragraph}", 800, 100)[0] == paragraph
    sentences = ("Una oración corta y completa. " * 40).strip()
    assert all(chunk.endswith(".") for chunk in split_text(sentences, 800, 100))


@pytest.mark.parametrize("size,overlap", [(0, 0), (100, 100), (100, -1)])
def test_invalid_chunk_parameters(size, overlap):
    with pytest.raises(ValueError):
        split_text("text", size, overlap)
