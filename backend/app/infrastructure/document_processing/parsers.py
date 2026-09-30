import asyncio
import re
import zipfile
from collections.abc import Iterator
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any

import docx
from pypdf import PdfReader

from app.application.ports import DocumentParser
from app.domain.exceptions import (
    DocumentTooLarge,
    InvalidDocument,
    UnsupportedDocumentType,
)

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_MIME_TYPES = {".pdf": "application/pdf", ".txt": "text/plain", ".docx": DOCX_MIME}
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
# Deleted or moved-away text, formatting, and the legacy copy of drawings
# (mc:Fallback duplicates text boxes) are not part of the reading text.
_SKIPPED = {
    f"{_W}del",
    f"{_W}moveFrom",
    f"{_W}pPr",
    f"{_W}rPr",
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback",
}


class TextParser:
    def parse(self, content: bytes, max_chars: int) -> str:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise InvalidDocument(
                "El archivo TXT debe usar codificación UTF-8."
            ) from exc
        if len(text) > max_chars:
            raise DocumentTooLarge()
        if any(ord(char) < 32 and char not in "\n\r\t\f" for char in text):
            raise InvalidDocument("El archivo TXT contiene datos binarios.")
        return text


class PDFParser:
    def __init__(self, max_pages: int = 200) -> None:
        self.max_pages = max_pages

    def parse(self, content: bytes, max_chars: int) -> str:
        try:
            reader = PdfReader(BytesIO(content))
            # Many PDFs only carry an owner password (print/copy restrictions)
            # and open without one; only a required user password is rejected.
            if reader.is_encrypted and not reader.decrypt(""):
                raise InvalidDocument("No se admiten PDF protegidos con contraseña.")
            if len(reader.pages) > self.max_pages:
                raise DocumentTooLarge()
            parts: list[str] = []
            length = 0
            for page in reader.pages:
                text = page.extract_text() or ""
                length += len(text) + 1
                if length > max_chars:
                    raise DocumentTooLarge()
                parts.append(text)
            return "\n".join(parts)
        except (InvalidDocument, DocumentTooLarge):
            raise
        except Exception as exc:
            raise InvalidDocument("No se pudo leer el PDF.") from exc


def _inline_text(element: Any) -> Iterator[str]:
    for child in element:
        if child.tag in _SKIPPED or child.tag == f"{_W}txbxContent":
            continue
        if child.tag == f"{_W}t":
            yield child.text or ""
        elif child.tag == f"{_W}tab":
            yield "\t"
        elif child.tag in {f"{_W}br", f"{_W}cr"}:
            yield "\n"
        else:
            yield from _inline_text(child)


def _text_boxes(element: Any) -> Iterator[Any]:
    for child in element:
        if child.tag in _SKIPPED:
            continue
        if child.tag == f"{_W}txbxContent":
            yield child
        else:
            yield from _text_boxes(child)


def _docx_lines(element: Any) -> Iterator[str]:
    """Paragraphs and table rows in reading order.

    Unlike python-docx's `paragraphs`, this includes tables, content controls,
    tracked insertions and text boxes; horizontally merged cells appear once.
    """
    for child in element:
        if child.tag in _SKIPPED:
            continue
        if child.tag == f"{_W}p":
            yield "".join(_inline_text(child))
            for box in _text_boxes(child):
                yield from _docx_lines(box)
        elif child.tag == f"{_W}tbl":
            for row in child.iterchildren(f"{_W}tr"):
                cells = (
                    " ".join(_docx_lines(cell)).strip()
                    for cell in row.iterchildren(f"{_W}tc")
                )
                yield " | ".join(cell for cell in cells if cell)
        else:
            yield from _docx_lines(child)


class DocxParser:
    def __init__(self, max_unzipped: int = 100 * 1024 * 1024) -> None:
        self.max_unzipped = max_unzipped

    def parse(self, content: bytes, max_chars: int) -> str:
        try:
            # A small archive can expand to gigabytes: check declared sizes first;
            # zipfile never reads past them.
            with zipfile.ZipFile(BytesIO(content)) as archive:
                unzipped = sum(item.file_size for item in archive.infolist())
            if unzipped > self.max_unzipped:
                raise DocumentTooLarge()
            body = docx.Document(BytesIO(content)).element.body
            lines: list[str] = []
            length = 0
            for line in _docx_lines(body):
                length += len(line) + 1
                if length > max_chars:
                    raise DocumentTooLarge()
                lines.append(line)
            return "\n".join(lines)
        except (InvalidDocument, DocumentTooLarge):
            raise
        except Exception as exc:
            raise InvalidDocument("No se pudo leer el documento de Word.") from exc


class FileDocumentProcessor:
    def __init__(
        self,
        max_bytes: int = 10 * 1024 * 1024,
        max_chars: int = 200_000,
        max_pages: int = 200,
    ) -> None:
        self.max_bytes = max_bytes
        self.max_chars = max_chars
        self.parsers: dict[str, DocumentParser] = {
            ".txt": TextParser(),
            ".pdf": PDFParser(max_pages),
            ".docx": DocxParser(10 * max_bytes),
        }

    def validate(self, filename: str, content_type: str, content: bytes) -> str:
        name = PurePosixPath(filename.replace("\\", "/")).name
        extension = PurePosixPath(name).suffix.lower()
        mime = content_type.split(";", 1)[0].lower().strip()
        if extension not in _MIME_TYPES or mime not in {
            _MIME_TYPES.get(extension),
            "application/octet-stream",
            "",
        }:
            raise UnsupportedDocumentType()
        if not content:
            raise InvalidDocument()
        if len(content) > self.max_bytes:
            raise DocumentTooLarge()
        if extension == ".pdf" and not content.startswith(b"%PDF-"):
            raise InvalidDocument("El contenido del archivo no corresponde a un PDF.")
        if extension == ".docx" and not content.startswith(b"PK\x03\x04"):
            # Password-protected Word files and legacy .doc use the OLE format.
            raise InvalidDocument(
                "No se admiten documentos de Word protegidos con contraseña "
                "ni en formato .doc."
                if content.startswith(b"\xd0\xcf\x11\xe0")
                else "El contenido del archivo no corresponde a un documento de Word."
            )
        stem = re.sub(r"[^\w. -]", "_", PurePosixPath(name).stem).strip(" .")
        return (stem or "documento")[:170] + extension

    async def extract(self, filename: str, content: bytes) -> str:
        parser = self.parsers[PurePosixPath(filename).suffix.lower()]
        text = await asyncio.to_thread(parser.parse, content, self.max_chars)
        # Keep line and paragraph breaks so titles do not merge into sentences.
        text = re.sub(r"[^\S\n]+", " ", text.replace("\x00", " "))
        text = re.sub(r"\n{3,}", "\n\n", re.sub(r" ?\n ?", "\n", text)).strip()
        if not text:
            raise InvalidDocument(
                "No se encontró texto. Los PDF escaneados y las imágenes requieren OCR."
            )
        return text
